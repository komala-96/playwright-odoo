"""
Generates Averigo_End_to_End_Test_Cases.xlsx from Averigo_End_to_End_Test_Plan_Complete.docx
Sheets: Test Cases | Test Data | Test Summary | Execution
"""
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

LOGIN = "Log in to the Averigo (Martiserv) QA application with a user account that has the module access named in Precondition."

# ---------------------------------------------------------------------------
# 5. Master Test Data (from Test Plan section 5)
# ---------------------------------------------------------------------------
MASTER_DATA = [
    ("Employee", "EMP-QA-001", "QA Test Employee"),
    ("Product Category", "CAT-QA-001", "QA Beverage"),
    ("UOM Type", "UOM-QA-001", "QA Case UoM"),
    ("Route", "ROUTE-QA-001", "QA Delivery Route"),
    ("Warehouse", "WH-QA-001", "QA Main Warehouse"),
    ("Truck", "TRUCK-QA-001", "QA Truck 01"),
    ("Bin", "BIN-QA-001", "QA-MAIN-BIN-01"),
    ("Parent Customer", "PARENT-QA-001", "QA Parent Customer"),
    ("Order Type", "ORDER-QA-001", "QA Standard Order"),
    ("Product 1", "PROD-QA-001", "Tax ON + Container Deposit ON"),
    ("Product 2", "PROD-QA-002", "Tax ON + Container Deposit OFF"),
    ("Product 3", "PROD-QA-003", "Tax OFF + Container Deposit ON"),
    ("Product 4", "PROD-QA-004", "Tax OFF + Container Deposit OFF"),
    ("Service Product", "SERV-QA-001", "QA Service Product"),
    ("Customer 1", "CUST-QA-001", "Customer with Fuel Charge, Hazard Fee & Subsidy"),
    ("Customer 2", "CUST-QA-002", "Customer without those taxes"),
    ("Vendor", "VEND-QA-001", "QA Test Vendor"),
    ("Micro Market", "MM-QA-001", "QA Micro Market 01"),
]

# ---------------------------------------------------------------------------
# 6. Transaction Test Data (from Test Plan section 6)
# ---------------------------------------------------------------------------
TXN_DATA = [
    ("Standard Sales Quantity", "10"),
    ("Large Order Quantity", "100"),
    ("Backorder / Partial Delivery", "50 delivered + 50 backordered"),
    ("Quantity Adjustment", "+20"),
    ("Value Adjustment", "+100.00"),
    ("Customer Advance", "500.00"),
    ("Vendor Advance", "500.00"),
    ("Discount", "10%"),
    ("S&H", "25.00"),
    ("Insurance", "15.00"),
    ("Payment", "Half of invoice/bill total, followed by remaining balance where applicable"),
    ("Product Return", "Half of invoiced quantity, then remaining quantity"),
]

# ---------------------------------------------------------------------------
# 7. Test Cases: (id, module, scenario, precondition, test_data, steps[list], expected, priority, type)
# ---------------------------------------------------------------------------
TC = []

def add(id_, module, scenario, precondition, data, steps, expected, priority, type_):
    TC.append({
        "id": id_, "module": module, "scenario": scenario,
        "precondition": precondition, "data": data,
        "steps": "\n".join(f"{i+1}. {s}" for i, s in enumerate([LOGIN] + steps)),
        "expected": expected, "priority": priority, "type": type_,
    })

# ---- Inventory / Customer Care master setup (TC-01..TC-11) ----
add("TC-01", "Inventory", "Create employee",
    "QA environment is available; user has Employee module access.",
    "EMP-QA-001 / QA Test Employee",
    ["Open the Employees app and click New.",
     "Enter Employee Name 'QA Test Employee', Employee Code EMP-QA-001, and mandatory fields (job position, department, work email/phone).",
     "Click Save.",
     "Open the Account Manager / Purchase Manager selection field on a customer or vendor form and confirm the employee is listed."],
    "Employee is created successfully and is available for account/purchase manager selection.",
    "High", "Functional")

add("TC-02", "Inventory", "Create product category, UOM type and route",
    "TC-01 completed; Inventory module access available.",
    "CAT-QA-001 / QA Beverage; UOM-QA-001 / QA Case UoM; ROUTE-QA-001 / QA Delivery Route",
    ["Open Inventory > Configuration > Product Categories, click New, enter code/name CAT-QA-001 'QA Beverage', Save.",
     "Open Inventory > Configuration > UoM Types, click New, enter code UOM-QA-001 'QA Case UoM' with reference unit and conversion factor, Save.",
     "Open Inventory > Configuration > Routes, click New, enter code/name ROUTE-QA-001 'QA Delivery Route', Save.",
     "Open the Product creation form and confirm all three masters are selectable."],
    "All masters are created successfully and are selectable in product setup.",
    "Medium", "Functional")

add("TC-03", "Inventory", "Create warehouse",
    "TC-01 completed; Inventory module access available.",
    "WH-QA-001 / QA Main Warehouse",
    ["Open Inventory > Locations/Warehouses, click New.",
     "Enter warehouse code/name WH-QA-001 'QA Main Warehouse' and mandatory address/company details.",
     "Click Save and confirm the warehouse appears under Inventory > Location."],
    "Warehouse is created under Inventory/Location and is available for stock operations.",
    "High", "Functional")

add("TC-04", "Inventory", "Create truck",
    "TC-03 completed; Inventory configuration access available.",
    "TRUCK-QA-001 / QA Truck 01",
    ["Open Inventory > Configuration > Trucks, click New.",
     "Enter truck code/name TRUCK-QA-001 'QA Truck 01', license plate, capacity and assigned warehouse WH-QA-001.",
     "Click Save."],
    "Truck is created successfully and is available for relevant delivery operations.",
    "Medium", "Functional")

add("TC-05", "Inventory", "Create bin location under new warehouse",
    "TC-03 completed; WH-QA-001 exists.",
    "BIN-QA-001 / QA-MAIN-BIN-01; WH-QA-001",
    ["Open Inventory > Bin Locations, click New.",
     "Enter bin code BIN-QA-001 'QA-MAIN-BIN-01' and select Warehouse WH-QA-001.",
     "Click Save and confirm the bin is listed under WH-QA-001."],
    "Bin is created and correctly associated with the new warehouse.",
    "High", "Functional")

add("TC-06", "Customer Care", "Create parent customer and order type",
    "TC-01 completed; Customer Care module access available.",
    "PARENT-QA-001 / QA Parent Customer; ORDER-QA-001 / QA Standard Order",
    ["Open Customer Care > Configuration > Parent Customers, click New, enter code/name PARENT-QA-001 'QA Parent Customer', Save.",
     "Open Customer Care > Configuration > Order Types, click New, enter code/name ORDER-QA-001 'QA Standard Order', Save.",
     "Open the Customer creation form and confirm both records are selectable."],
    "Both configuration records are created and selectable.",
    "Medium", "Functional")

add("TC-07", "Inventory", "Create Product 1: Tax ON + Deposit ON",
    "TC-02, TC-05 completed.",
    "PROD-QA-001 / Tax ON + Container Deposit ON; CAT-QA-001; UOM-QA-001; BIN-QA-001",
    ["Open Inventory > Products, click New.",
     "Enter Product Name/Code PROD-QA-001, select Category CAT-QA-001, UOM Type UOM-QA-001, enter a UPC code and set Primary Bin to BIN-QA-001.",
     "Enable the Tax checkbox/field and enable Container Deposit with a deposit amount.",
     "Click Save."],
    "Product is created with required UOM, category, UPC and primary bin; tax and container deposit are enabled.",
    "High", "Functional")

add("TC-08", "Inventory", "Create Product 2: Tax ON + Deposit OFF",
    "TC-02, TC-05 completed.",
    "PROD-QA-002 / Tax ON + Container Deposit OFF",
    ["Open Inventory > Products, click New.",
     "Enter Product Name/Code PROD-QA-002 with the same category/UOM/bin masters used in TC-07.",
     "Enable Tax and leave Container Deposit disabled.",
     "Click Save."],
    "Product is created with required masters and tax enabled/deposit disabled.",
    "High", "Functional")

add("TC-09", "Inventory", "Create Product 3: Tax OFF + Deposit ON",
    "TC-02, TC-05 completed.",
    "PROD-QA-003 / Tax OFF + Container Deposit ON",
    ["Open Inventory > Products, click New.",
     "Enter Product Name/Code PROD-QA-003 with the same category/UOM/bin masters used in TC-07.",
     "Disable Tax and enable Container Deposit with a deposit amount.",
     "Click Save."],
    "Product is created with required masters and tax disabled/deposit enabled.",
    "High", "Functional")

add("TC-10", "Inventory", "Create Product 4: Tax OFF + Deposit OFF",
    "TC-02, TC-05 completed.",
    "PROD-QA-004 / Tax OFF + Container Deposit OFF",
    ["Open Inventory > Products, click New.",
     "Enter Product Name/Code PROD-QA-004 with the same category/UOM/bin masters used in TC-07.",
     "Disable both Tax and Container Deposit.",
     "Click Save."],
    "Product is created with required masters and both tax and deposit disabled.",
    "High", "Functional")

add("TC-11", "Inventory", "Create service product",
    "TC-02 completed.",
    "SERV-QA-001 / QA Service Product",
    ["Open Inventory > Products, click New and set Product Type to Service.",
     "Enter Product Name/Code SERV-QA-001 'QA Service Product' and required sales price/category fields.",
     "Click Save.",
     "Open Inventory On-Hand / Stock Valuation and confirm no stock quantity record is created for this product."],
    "Service product is created successfully and does not create physical inventory movement.",
    "Medium", "Functional")

# ---- Reports on masters (TC-12, 13) ----
add("TC-12", "Reports", "Run product addition/deletion report",
    "TC-07 to TC-11 completed.",
    "PROD-QA-001..004, SERV-QA-001",
    ["Open Reports > Inventory > Product Addition/Deletion Report.",
     "Filter by the date range covering today's product creation activity and run the report.",
     "Verify each of the 5 newly created products appears as an addition with the correct creation date/user."],
    "Created product additions/deletions are displayed accurately.",
    "Medium", "Report")

add("TC-13", "Reports", "Run product master report",
    "TC-07 to TC-11 completed.",
    "PROD-QA-001..004, SERV-QA-001",
    ["Open Reports > Inventory > Product Master Report.",
     "Run the report without filters (or filter by category CAT-QA-001).",
     "Verify all four products and the service product appear with correct category, UOM, tax and container-deposit details."],
    "All four products and service product appear with correct master details.",
    "Medium", "Report")

# ---- Adjustments & related reports (TC-14..16) ----
add("TC-14", "Inventory", "Perform quantity and value adjustment on Product 1",
    "TC-07 completed; PROD-QA-001 has existing stock in WH-QA-001.",
    "Quantity Adjustment +20; Value Adjustment +100.00",
    ["Open Inventory > Inventory Adjustments, click New, select WH-QA-001 and PROD-QA-001, enter a Quantity Adjustment reference and increase quantity by +20.",
     "Click Validate Inventory to post the quantity adjustment.",
     "Create a second Inventory Adjustment record for a Value Adjustment on PROD-QA-001 with a new cost reflecting +100.00, and Validate Inventory.",
     "Open the product's On Hand quantities and confirm the updated stock and cost."],
    "Quantity/value adjustments are posted and resulting stock is correct.",
    "High", "Functional")

add("TC-15", "Reports", "Run warehouse stock value report",
    "TC-14 completed.",
    "WH-QA-001; PROD-QA-001",
    ["Open Reports > Inventory > Stock Value Report (Warehouse).",
     "Filter by warehouse WH-QA-001 and run the report.",
     "Verify PROD-QA-001's quantity and value reflect the +20 quantity and +100.00 cost adjustments from TC-14."],
    "Stock quantity/value reflects the latest inventory transactions and adjustments.",
    "Medium", "Report")

add("TC-16", "Reports", "Run product cost change report",
    "TC-14 completed.",
    "PROD-QA-001",
    ["Open Reports > Inventory > Product Cost Change Report.",
     "Filter by product PROD-QA-001 and the date of the TC-14 value adjustment.",
     "Verify the report shows the correct before/after cost values."],
    "Product cost changes are captured with correct before/after information.",
    "Medium", "Report")

# ---- Customer catalogue and customers (TC-17..21) ----
add("TC-17", "Customer Care", "Create customer product catalogue",
    "TC-07 to TC-11 completed.",
    "Catalogue type: Customer; PROD-QA-002, PROD-QA-003, PROD-QA-004, SERV-QA-001",
    ["Open Inventory/Customer Care > Product Catalogue, click New and set Catalogue Type to 'Customer'.",
     "Add PROD-QA-002, PROD-QA-003, PROD-QA-004 and SERV-QA-001 to the catalogue; do NOT add PROD-QA-001.",
     "Set price/min/max quantity for each line and click Save."],
    "Catalogue of type Customer contains all created products except Product 1.",
    "Medium", "Functional")

add("TC-18", "Customer Care", "Create Customer 1 with taxes",
    "TC-06, TC-07, TC-17 completed.",
    "CUST-QA-001; PROD-QA-001 (Each UOM/new UOM); EMP-QA-001; Fuel Charge, Hazard Fee, Subsidy taxes",
    ["Open Customer Care, click New, enter Customer Name/Code CUST-QA-001, select Parent Customer PARENT-QA-001.",
     "Add PROD-QA-001 as an approved product using the Each UOM and the new UOM created in TC-02.",
     "Add the catalogue from TC-17 to bring in the remaining approved products.",
     "Enable/configure Fuel Charge, Hazard Fee and Subsidy taxes on the customer.",
     "Set EMP-QA-001 as the Account Manager.",
     "Click Save."],
    "Customer uses Product 1 with Each UOM/new UOM from catalogue; Fuel Charge, Hazard Fee and Subsidy are included; employee is account manager.",
    "High", "Functional")

add("TC-19", "Customer Care", "Create Customer 2 without taxes",
    "TC-06, TC-17 completed.",
    "CUST-QA-002",
    ["Open Customer Care, click New, enter Customer Name/Code CUST-QA-002, select Parent Customer PARENT-QA-001.",
     "Add the catalogue from TC-17 to bring in approved products.",
     "Leave Fuel Charge, Hazard Fee and Subsidy taxes disabled/unconfigured.",
     "Click Save."],
    "Customer is created without the specified Fuel Charge, Hazard Fee and Subsidy taxes.",
    "High", "Functional")

add("TC-20", "Reports", "Run customer detailed report",
    "TC-18, TC-19 completed.",
    "CUST-QA-001, CUST-QA-002",
    ["Open Reports > Customer Care > Customer Detailed Report.",
     "Filter by CUST-QA-001 and CUST-QA-002 (or Parent Customer PARENT-QA-001) and run the report.",
     "Verify master, tax configuration and account-manager (EMP-QA-001) details are correct for both customers."],
    "Both customers display correct master, tax and account-manager information.",
    "Medium", "Report")

add("TC-21", "Reports", "Run customer product list report",
    "TC-18, TC-19 completed.",
    "CUST-QA-001, CUST-QA-002",
    ["Open Reports > Customer Care > Customer Product List Report.",
     "Filter by CUST-QA-001 and CUST-QA-002 and run the report.",
     "Verify the approved product / UOM associations match the products added in TC-18 and TC-19."],
    "Customer-product associations and UOM details are correct.",
    "Medium", "Report")

# ---- Sales / AR flows (TC-22..49) ----
add("TC-22", "Sales", "Customer 1 order + invoice with service product and half payment",
    "TC-18 completed.",
    "CUST-QA-001; PROD-QA-001 Qty 10; SERV-QA-001; Discount 10%; S&H 25.00; Insurance 15.00; Payment: half of invoice total",
    ["Open Customer Care > Orders, click New, select CUST-QA-001.",
     "Add PROD-QA-001 with Quantity 10 and add the SERV-QA-001 service line.",
     "Apply Discount 10%, S&H 25.00 and Insurance 15.00, then Save and note the order's Subtotal/Total.",
     "Create Invoice from the order and note the invoice's Subtotal, S&H, Insurance, Discount and Total.",
     "Compare order and invoice totals for a match.",
     "Click Register Payment, enter an amount equal to half the invoice Total, select Mode of Payment and Deposit To account, and Save."],
    "Order and invoice Subtotal, S&H, Insurance, Discount and Total match; service line is included; half payment reduces due balance correctly.",
    "High", "Financial")

add("TC-23", "AR", "AR payment receipt for same invoice",
    "TC-22 completed; invoice has a remaining balance.",
    "CUST-QA-001; Payment: remaining balance",
    ["Open Accounting/AR > Payments > Customer Payment Receipt, click New.",
     "Select Customer CUST-QA-001 and the invoice from TC-22, and apply the payment against the remaining due amount.",
     "Click Save/Post and note the Receipt Number.",
     "Open the invoice and confirm the balance due is now zero."],
    "AR payment receipt is posted against the same invoice and customer balance updates correctly.",
    "High", "Financial")

add("TC-24", "Sales", "Customer 1 order + invoice without service product",
    "TC-18 completed.",
    "CUST-QA-001; PROD-QA-001 Qty 10; Discount 10%; S&H 25.00; Insurance 15.00",
    ["Open Customer Care > Orders, click New, select CUST-QA-001.",
     "Add PROD-QA-001 with Quantity 10 only (no service line).",
     "Apply Discount 10%, S&H 25.00 and Insurance 15.00, Save and note order totals.",
     "Create Invoice and compare Subtotal, S&H, Insurance, Discount and Total against the order."],
    "Order and invoice financial values match and no service product is included.",
    "High", "Financial")

add("TC-25", "Sales/Delivery", "Customer 2 order for Product 1 + packing slip + delivery management",
    "TC-19 completed.",
    "CUST-QA-002; PROD-QA-001 Qty 10",
    ["Open Customer Care > Orders, click New, select CUST-QA-002 and add PROD-QA-001 with Quantity 10.",
     "Save the order and generate a Packing Slip for the order.",
     "Open Delivery Management, select the packing slip, and confirm/complete the delivery.",
     "Create Invoice from the delivered order and compare order vs. invoice Subtotal, S&H, Insurance, Discount and Total (including service product if one was selected)."],
    "Packing slip/delivery completes successfully; order and invoice financial values match, including service product where selected.",
    "High", "Integration")

add("TC-26", "Sales", "Customer 2 Product 1 invoice without service product",
    "TC-25 completed.",
    "CUST-QA-002; PROD-QA-001",
    ["Open the Customer 2 order for PROD-QA-001 (no service line) and create/verify its invoice.",
     "Compare order and invoice Subtotal, S&H, Insurance, Discount and Total."],
    "Subtotal, S&H, Insurance, Discount and Total match between order and invoice.",
    "Medium", "Financial")

add("TC-27", "Sales", "Customer 2 Product 2 invoice with service product",
    "TC-19 completed.",
    "CUST-QA-002; PROD-QA-002; SERV-QA-001",
    ["Open Customer Care > Orders, click New, select CUST-QA-002, add PROD-QA-002 and the SERV-QA-001 service line.",
     "Save, create Invoice and confirm order/invoice values match and the service line is present."],
    "Order/invoice values match and service product is included.",
    "Medium", "Financial")

add("TC-28", "Sales", "Customer 2 Product 2 invoice without service product",
    "TC-19 completed.",
    "CUST-QA-002; PROD-QA-002",
    ["Open Customer Care > Orders, click New, select CUST-QA-002 and add PROD-QA-002 only.",
     "Save, create Invoice and confirm order/invoice values match."],
    "Order/invoice values match without service product.",
    "Medium", "Financial")

add("TC-29", "Sales", "Customer 2 Product 3 invoice with service product",
    "TC-19 completed.",
    "CUST-QA-002; PROD-QA-003; SERV-QA-001",
    ["Open Customer Care > Orders, click New, select CUST-QA-002, add PROD-QA-003 and the SERV-QA-001 service line.",
     "Save, create Invoice and confirm order/invoice values match and the service line is present."],
    "Order/invoice values match and service product is included.",
    "Medium", "Financial")

add("TC-30", "Sales", "Customer 2 Product 3 invoice without service product",
    "TC-19 completed.",
    "CUST-QA-002; PROD-QA-003",
    ["Open Customer Care > Orders, click New, select CUST-QA-002 and add PROD-QA-003 only.",
     "Save, create Invoice and confirm order/invoice values match."],
    "Order/invoice values match without service product.",
    "Medium", "Financial")

add("TC-31", "Sales", "Customer 2 Product 4 invoice with service product",
    "TC-19 completed.",
    "CUST-QA-002; PROD-QA-004; SERV-QA-001",
    ["Open Customer Care > Orders, click New, select CUST-QA-002, add PROD-QA-004 and the SERV-QA-001 service line.",
     "Save, create Invoice and confirm order/invoice values match and the service line is present."],
    "Order/invoice values match and service product is included.",
    "Medium", "Financial")

add("TC-32", "Sales", "Customer 2 Product 4 invoice without service product",
    "TC-19 completed.",
    "CUST-QA-002; PROD-QA-004",
    ["Open Customer Care > Orders, click New, select CUST-QA-002 and add PROD-QA-004 only.",
     "Save, create Invoice and confirm order/invoice values match."],
    "Order/invoice values match without service product.",
    "Medium", "Financial")

add("TC-33", "Sales", "Create order for service product",
    "TC-11 completed.",
    "CUST-QA-002; SERV-QA-001",
    ["Open Customer Care > Orders, click New, select CUST-QA-002 and add SERV-QA-001 as the only line.",
     "Save the order and create its invoice.",
     "Open Inventory On-Hand for SERV-QA-001 and confirm no physical stock quantity changed."],
    "Service-only order is accepted and does not incorrectly affect physical stock.",
    "Medium", "Functional")

add("TC-34", "Reports", "Run product price change report",
    "TC-07 to TC-10 completed and at least one product price edited.",
    "PROD-QA-001..004",
    ["Edit the sales price of PROD-QA-001 and Save.",
     "Open Reports > Inventory > Product Price Change Report.",
     "Filter by PROD-QA-001 and verify the before/after price is reported accurately."],
    "Price changes are reported accurately.",
    "Medium", "Report")

add("TC-35", "Reports", "Run product UOM change report",
    "TC-07 to TC-10 completed and a product UOM edited.",
    "PROD-QA-001..004",
    ["Edit the UOM configuration of PROD-QA-001 (e.g. add/change a case/box UOM) and Save.",
     "Open Reports > Inventory > Product UOM Change Report.",
     "Filter by PROD-QA-001 and verify the UOM change is reported accurately."],
    "UOM changes are reported accurately.",
    "Medium", "Report")

add("TC-36", "AR", "Create customer advance",
    "TC-18 completed.",
    "CUST-QA-001; Customer Advance 500.00",
    ["Open Accounting/AR > Customer Advance Payment, click New.",
     "Select Customer CUST-QA-001, enter Amount 500.00, select Mode of Payment/Deposit To account.",
     "Click Save/Post and note the advance reference.",
     "Open the customer's account balance and confirm it reflects the 500.00 advance."],
    "Advance is posted and customer balance reflects the advance correctly.",
    "High", "Financial")

add("TC-37", "Sales Return", "Return half of invoiced quantity and create credit memo",
    "TC-24 completed (Customer 1, Qty 10 invoice); TC-36 advance available.",
    "CUST-QA-001; PROD-QA-001; Return Qty 5 (half of 10)",
    ["Open Customer Care > Returns, click New, select the CUST-QA-001 invoice from TC-24 and PROD-QA-001.",
     "Enter Return Quantity 5 (half of the invoiced Qty 10).",
     "Click Save/Validate to post the return and confirm a Credit Memo is generated.",
     "Verify the returned stock is added back to WH-QA-001 and the customer/advance balance reflects the credit."],
    "Returned stock and customer/advance amount are updated correctly; credit memo is created.",
    "High", "Financial")

add("TC-38", "Sales Return", "Return remaining quantity of same invoice",
    "TC-37 completed.",
    "CUST-QA-001; PROD-QA-001; Return Qty 5 (remaining)",
    ["Open Customer Care > Returns, click New against the same TC-24 invoice and PROD-QA-001.",
     "Enter Return Quantity 5 (remaining balance of the original Qty 10).",
     "Click Save/Validate and confirm the credit memo is generated.",
     "Verify cumulative returned quantity for the invoice now equals 10 and stock/credit balances reconcile."],
    "Remaining return is processed correctly and cumulative returned quantity is accurate.",
    "High", "Financial")

add("TC-39", "Sales Return", "Create all three credit memo types",
    "TC-37, TC-38 completed.",
    "CUST-QA-001",
    ["Open Customer Care > Credit Memo, click New and create a Credit Memo of the first supported type (e.g. Product Return) referencing TC-37/TC-38.",
     "Create a second Credit Memo of the second supported type (e.g. Price Adjustment) for CUST-QA-001.",
     "Create a third Credit Memo of the third supported type (e.g. Direct/Manual Credit) for CUST-QA-001.",
     "Verify each credit memo posts with the correct accounting impact on the customer balance."],
    "Each supported credit memo type is created successfully with correct accounting impact.",
    "Medium", "Financial")

add("TC-40", "AR", "Create direct invoice for Customer 1 and pay",
    "TC-18 completed.",
    "CUST-QA-001; Payment: full amount",
    ["Open Accounting/AR > Direct Invoice, click New, select CUST-QA-001, add a line item and Save/Post.",
     "Click Register Payment, enter the full invoice amount, select Mode of Payment and Deposit To account, and Save.",
     "Verify the invoice balance is fully cleared."],
    "Direct invoice is created and payment clears/reduces the balance correctly.",
    "Medium", "Financial")

add("TC-41", "AR", "Create direct invoice for Customer 2 with Product 1 and pay",
    "TC-19 completed.",
    "CUST-QA-002; PROD-QA-001; Payment: full amount",
    ["Open Accounting/AR > Direct Invoice, click New, select CUST-QA-002 and add PROD-QA-001 as the line item.",
     "Save/Post the invoice and verify the stock/financial impact on PROD-QA-001.",
     "Register Payment for the full invoice amount and Save."],
    "Product line, stock/financial impact and payment are correct.",
    "Medium", "Financial")

add("TC-42", "AR", "Create direct invoice for service product",
    "TC-11 completed.",
    "CUST-QA-002; SERV-QA-001",
    ["Open Accounting/AR > Direct Invoice, click New, select CUST-QA-002 and add SERV-QA-001 as the line item.",
     "Save/Post the invoice.",
     "Open Inventory On-Hand for SERV-QA-001 and confirm no physical stock movement occurred."],
    "Service invoice is created correctly without physical stock movement.",
    "Medium", "Financial")

add("TC-43", "Reports", "Run AR/customer reports",
    "TC-22 to TC-42 completed.",
    "CUST-QA-001, CUST-QA-002",
    ["Open Reports > AR > Aging by Customer, filter by CUST-QA-001/CUST-QA-002, run and verify aging buckets.",
     "Open Customer Statement report for both customers and verify transaction history/balances.",
     "Open Invoice List and Invoice Export reports and verify all invoices from TC-22 to TC-42 are listed with correct amounts.",
     "Open Receipt List report and verify payments/receipts posted in TC-22, TC-23, TC-40, TC-41 are listed.",
     "Open Revenue by Customer and Revenue by Customer (Delivery) reports and verify revenue totals match posted invoices."],
    "Aging by Customer, Customer Statement, Invoice List, Invoice Export, Receipt List, Revenue by Customer and Revenue by Customer (Delivery) show correct values.",
    "Medium", "Report")

add("TC-44", "Sales/Backorder", "Create 100-quantity order, reduce to 50, create backorder and deliver",
    "TC-18 or TC-19 completed.",
    "CUST-QA-002; PROD-QA-001; Large Order Qty 100; Backorder 50 delivered + 50 backordered",
    ["Open Customer Care > Orders, click New, select CUST-QA-002, add PROD-QA-001 with Quantity 100 and Save.",
     "On the Delivery tab, reduce the delivered quantity to 50, confirming the remaining 50 is placed on backorder.",
     "Create Invoice for the delivered 50 quantity and note the invoice number/amount.",
     "Process the backorder for the remaining 50 quantity, complete delivery and create the second invoice.",
     "Verify two invoices exist totalling Quantity 100."],
    "Outstanding quantity is backordered correctly; after delivery, two invoices are created as expected.",
    "High", "Integration")

add("TC-45", "Reports", "Run Where Sold (Warehouse) report",
    "TC-22 to TC-44 completed.",
    "PROD-QA-001; WH-QA-001",
    ["Open Reports > Sales > Where Sold (Warehouse) Report.",
     "Filter by PROD-QA-001 and/or WH-QA-001 and run the report.",
     "Verify the customers/orders selling PROD-QA-001 from WH-QA-001 are correctly attributed."],
    "Product sales are correctly attributed to the warehouse.",
    "Medium", "Report")

add("TC-46", "Reports", "Run Customer Order Activity report",
    "TC-22 to TC-44 completed.",
    "CUST-QA-001, CUST-QA-002",
    ["Open Reports > Sales > Customer Order Activity Report.",
     "Filter by CUST-QA-001 and CUST-QA-002 and run the report.",
     "Verify all orders created in TC-22 to TC-44 are listed with correct status/values."],
    "Customer order activity is accurate.",
    "Medium", "Report")

add("TC-47", "Reports", "Run Order List report",
    "TC-22 to TC-44 completed.",
    "CUST-QA-001, CUST-QA-002",
    ["Open Reports > Sales > Order List Report.",
     "Filter by the QA date range/customers and run the report.",
     "Verify orders and their statuses (Confirmed, Backordered, Delivered) match the transactions performed."],
    "Orders and statuses are accurately listed.",
    "Medium", "Report")

add("TC-48", "Reports", "Run Order Margin report",
    "TC-22 to TC-44 completed.",
    "PROD-QA-001..004",
    ["Open Reports > Sales > Order Margin Report.",
     "Filter by the QA orders/products and run the report.",
     "Verify margin (sales price vs. cost) calculations reconcile with the underlying order and cost data."],
    "Margin calculations reflect sales/cost data accurately.",
    "Medium", "Report")

add("TC-49", "Reports", "Run Sales Tax report",
    "TC-22, TC-24, TC-26 to TC-32 completed (Customer 1 and Customer 2 orders).",
    "CUST-QA-001 (taxes ON), CUST-QA-002 (taxes OFF)",
    ["Open Reports > Sales > Sales Tax Report.",
     "Filter by CUST-QA-001 and CUST-QA-002 and run the report.",
     "Verify Fuel Charge/Hazard Fee/Subsidy tax values appear for CUST-QA-001 and are absent for CUST-QA-002, matching each customer's tax configuration."],
    "Tax values reflect the applicable product/customer tax configuration.",
    "Medium", "Report")

# ---- Deletion checks (TC-50, TC-51) ----
add("TC-50", "Deletion", "Delete product and customer",
    "TC-07 (or any transacted product) and TC-18/TC-19 (transacted customer) completed.",
    "PROD-QA-001 (used in transactions); CUST-QA-001 (used in transactions)",
    ["Open Inventory > Products, select PROD-QA-001 (used in TC-22/TC-24/etc.) and attempt to delete it.",
     "Confirm the system displays an alert/restriction message and blocks the deletion due to existing dependencies.",
     "Open Customer Care, select CUST-QA-001 and attempt to delete it.",
     "Confirm the system displays the equivalent alert/restriction message and blocks the deletion."],
    "System displays the expected alert/restriction message and prevents unsafe deletion where dependencies exist.",
    "High", "Regression")

add("TC-51", "Reports", "Rerun product addition/deletion report after deletion attempt",
    "TC-50 completed.",
    "PROD-QA-001",
    ["Open Reports > Inventory > Product Addition/Deletion Report and re-run it for the current date.",
     "Verify PROD-QA-001 still shows only its original addition entry and no false deletion entry was introduced by the blocked TC-50 attempt."],
    "Report reflects actual product lifecycle activity and does not falsely show an invalid deletion.",
    "Medium", "Report")

# ---- Purchase / AP (TC-52..70) ----
add("TC-52", "Purchase", "Create vendor",
    "TC-01, TC-07 completed.",
    "VEND-QA-001; PROD-QA-001 (Each/new UOM); EMP-QA-001",
    ["Open Purchase > Vendors, click New, enter Vendor Name/Code VEND-QA-001.",
     "Add PROD-QA-001 as a supplied product using the Each UOM and the new UOM created in TC-02.",
     "Set EMP-QA-001 as the Purchase Manager.",
     "Click Save."],
    "Vendor is created with Product 1, Each/new UOM and employee as purchase manager.",
    "High", "Functional")

add("TC-53", "Purchase", "PO + bill for Product 1",
    "TC-52 completed.",
    "VEND-QA-001; PROD-QA-001 Qty 10; Discount 10%; S&H 25.00; Insurance 15.00; Payment: half then remaining",
    ["Open Purchase > Purchase Orders, click New, select VEND-QA-001 and add PROD-QA-001 with Quantity 10.",
     "Apply Discount 10%, S&H 25.00 and Insurance 15.00, Save and note the PO Subtotal/Total.",
     "Confirm the PO and receive the stock into WH-QA-001.",
     "Create the Bill from the PO and compare Subtotal, S&H, Insurance, Discount and Total against the PO.",
     "Register a payment for half the bill total, then a second payment for the remaining balance."],
    "PO and bill Subtotal, S&H, Insurance, Discount and Total match; stock updates and bill payment posts correctly.",
    "High", "Financial")

add("TC-54", "Purchase", "PO + bill for Product 2",
    "TC-52 completed.",
    "VEND-QA-001; PROD-QA-002",
    ["Open Purchase > Purchase Orders, click New, select VEND-QA-001 and add PROD-QA-002.",
     "Confirm the PO, receive stock, and create the Bill.",
     "Compare PO vs. bill financial values and confirm PROD-QA-002 stock updates correctly."],
    "Financial values match and Product 2 stock updates correctly.",
    "Medium", "Financial")

add("TC-55", "Purchase", "PO + bill for Product 3",
    "TC-52 completed.",
    "VEND-QA-001; PROD-QA-003",
    ["Open Purchase > Purchase Orders, click New, select VEND-QA-001 and add PROD-QA-003.",
     "Confirm the PO, receive stock, and create the Bill.",
     "Compare PO vs. bill financial values and confirm PROD-QA-003 stock updates correctly."],
    "Financial values match and Product 3 stock updates correctly.",
    "Medium", "Financial")

add("TC-56", "Purchase", "PO + bill for Product 4",
    "TC-52 completed.",
    "VEND-QA-001; PROD-QA-004",
    ["Open Purchase > Purchase Orders, click New, select VEND-QA-001 and add PROD-QA-004.",
     "Confirm the PO, receive stock, and create the Bill.",
     "Compare PO vs. bill financial values and confirm PROD-QA-004 stock updates correctly."],
    "Financial values match and Product 4 stock updates correctly.",
    "Medium", "Financial")

add("TC-57", "Purchase/AP", "Material issue and debit memo",
    "TC-53 completed (PROD-QA-001 received into stock).",
    "VEND-QA-001; PROD-QA-001",
    ["Open Inventory/Purchase > Material Issue, click New, select PROD-QA-001 and issue a quantity against VEND-QA-001.",
     "Save/Validate and confirm the stock is reduced.",
     "Open Purchase/AP > Debit Memo, create a debit memo referencing the material issue for VEND-QA-001.",
     "Verify the vendor advance/balance reflects the debit memo."],
    "Material issue updates stock; debit memo is created and vendor advance/balance impact is correct.",
    "High", "Financial")

add("TC-58", "Purchase/AP", "Vendor advance and all three debit memo types",
    "TC-52 completed.",
    "VEND-QA-001; Vendor Advance 500.00",
    ["Open Accounting/AP > Vendor Advance Payment, click New, select VEND-QA-001, enter Amount 500.00, Save/Post.",
     "Verify the vendor balance reflects the 500.00 advance.",
     "Create a Debit Memo of the first supported type for VEND-QA-001.",
     "Create a Debit Memo of the second supported type for VEND-QA-001.",
     "Create a Debit Memo of the third supported type for VEND-QA-001 and verify each posts with correct accounting impact."],
    "Vendor advance updates balance correctly and each supported debit memo type is created successfully.",
    "Medium", "Financial")

add("TC-59", "AP", "Direct bill with half payment and remaining payment",
    "TC-52 completed.",
    "VEND-QA-001; Payment: half then remaining",
    ["Open Accounting/AP > Direct Bill, click New, select VEND-QA-001, add a line item and Save/Post.",
     "Register a payment equal to half the bill total and Save; verify the due amount is reduced by half.",
     "Register a second payment for the remaining balance and verify the bill is fully cleared."],
    "First payment reduces due amount by half; remaining payment clears the invoice balance.",
    "Medium", "Financial")

add("TC-60", "Purchase", "Create direct receipt",
    "TC-52 completed.",
    "VEND-QA-001; PROD-QA-001",
    ["Open Purchase > Direct Receipt, click New, select VEND-QA-001 and PROD-QA-001 with a receipt quantity.",
     "Click Save/Validate to post the receipt.",
     "Verify PROD-QA-001 stock in WH-QA-001 increases and the accounting impact is correct."],
    "Direct receipt is created successfully and stock/accounting impact is correct.",
    "Medium", "Functional")

add("TC-61", "Purchase", "Create purchase backorder",
    "TC-52 completed.",
    "VEND-QA-001; PROD-QA-001; Large Order Qty 100; Backorder 50 + 50",
    ["Open Purchase > Purchase Orders, click New, select VEND-QA-001, add PROD-QA-001 with Quantity 100.",
     "Receive only 50 units, leaving the remaining 50 as an outstanding backordered quantity.",
     "Verify the backorder record retains the outstanding Quantity 50.",
     "Later, receive the remaining 50 units against the same backorder and confirm it clears."],
    "Backordered purchase quantity is retained accurately and can be received later.",
    "Medium", "Integration")

add("TC-62", "Purchase", "Delete vendor",
    "TC-52 to TC-61 completed (vendor has transactions).",
    "VEND-QA-001",
    ["Open Purchase > Vendors, select VEND-QA-001 and attempt to delete it.",
     "Confirm the system displays an alert/restriction message and blocks the deletion due to existing PO/bill dependencies."],
    "System displays the expected alert/restriction message when dependencies exist.",
    "High", "Regression")

add("TC-63", "Reports", "Run vendor invoice report",
    "TC-53 to TC-59 completed.",
    "VEND-QA-001",
    ["Open Reports > Purchase > Vendor Invoice Report.",
     "Filter by VEND-QA-001 and run the report.",
     "Verify all bills/direct bills from TC-53, TC-56, TC-59 are listed with correct amounts."],
    "Vendor invoice information and amounts are accurate.",
    "Medium", "Report")

add("TC-64", "Reports", "Run purchase order report",
    "TC-53 to TC-61 completed.",
    "VEND-QA-001",
    ["Open Reports > Purchase > Purchase Order Report.",
     "Filter by VEND-QA-001 and run the report.",
     "Verify all POs from TC-53 to TC-61 are listed with correct status and totals."],
    "Purchase orders are accurately listed.",
    "Medium", "Report")

add("TC-65", "Reports", "Run vendor list report",
    "TC-52 completed.",
    "VEND-QA-001",
    ["Open Reports > Purchase > Vendor List Report.",
     "Run the report and verify VEND-QA-001 appears with correct master data (purchase manager, supplied products)."],
    "Vendor master data is accurately listed.",
    "Low", "Report")

add("TC-66", "Reports", "Run inventory receipt report",
    "TC-53 to TC-61 completed.",
    "WH-QA-001; PROD-QA-001..004",
    ["Open Reports > Inventory > Inventory Receipt Report.",
     "Filter by WH-QA-001 and run the report.",
     "Verify receipts, products, quantities and source (PO/direct receipt) details from TC-53 to TC-61 are accurate."],
    "Receipts, products, quantities and source details are accurate.",
    "Medium", "Report")

add("TC-67", "Reports", "Run product order report",
    "TC-53 to TC-61 completed.",
    "PROD-QA-001..004",
    ["Open Reports > Purchase > Product Order Report.",
     "Filter by PROD-QA-001 to PROD-QA-004 and run the report.",
     "Verify order/quantity information is accurate for each product."],
    "Product order information is accurate.",
    "Medium", "Report")

add("TC-68", "Reports", "Run product transaction report (Inventory)",
    "TC-14, TC-53 to TC-61, TC-69, TC-70 completed.",
    "PROD-QA-001..004",
    ["Open Reports > Inventory > Product Transaction Report.",
     "Filter by PROD-QA-001 (or all QA products) and run the report.",
     "Verify all inventory transactions, quantities and movement types (adjustment, receipt, sale, transfer) are accurate."],
    "Inventory transactions, quantities and movement types are accurate.",
    "Medium", "Report")

add("TC-69", "Inventory", "Material transfer for Product 1",
    "TC-53 completed (PROD-QA-001 in WH-QA-001).",
    "PROD-QA-001; WH-QA-001; Quantity Adjustment reference qty",
    ["Open Inventory > Material Transfer, click New, select source WH-QA-001 and a destination warehouse/bin.",
     "Select PROD-QA-001 and enter a transfer quantity, then Save/Validate.",
     "Verify the source location stock decreases and the destination location stock increases by the transferred quantity."],
    "Stock is transferred correctly and source/destination quantities reconcile.",
    "Medium", "Functional")

add("TC-70", "Inventory", "Virtual location stock transfer",
    "TC-69 completed.",
    "PROD-QA-001; WH-QA-001",
    ["Open Inventory > Virtual Location Stock Transfer, click New, select PROD-QA-001, source WH-QA-001 and a virtual location.",
     "Enter the transfer quantity and Save/Validate.",
     "Verify source and destination (virtual location) stock quantities reconcile correctly."],
    "Virtual-location transfer updates source and destination stock correctly.",
    "Medium", "Functional")

# ---- Service Management / Equipment / Case (TC-71..80) ----
add("TC-71", "Service Management", "Create case categories, stages, origin and type",
    "QA environment available; Service Management configuration access.",
    "QA case configuration values",
    ["Open Service Management > Configuration > Case Categories, click New, create a QA category, Save.",
     "Open Configuration > Case Stages, create a QA stage, Save.",
     "Open Configuration > Case Origin, create a QA origin, Save.",
     "Open Configuration > Case Type, create a QA case type, Save.",
     "Confirm all four masters are selectable when creating a case."],
    "All service-management configuration masters are created successfully.",
    "Medium", "Functional")

add("TC-72", "Equipment", "Create equipment type, name, model, model no. and area/POS location",
    "QA environment available; Equipment module access.",
    "QA equipment configuration values",
    ["Open Equipment > Configuration > Equipment Type, click New, create a QA equipment type, Save.",
     "Open Configuration > Equipment Name, create a QA equipment name, Save.",
     "Open Configuration > Equipment Model, create a QA model, Save.",
     "Open Configuration > Equipment Model No., create a QA model number, Save.",
     "Open Configuration > Area/POS Location, create a QA area/POS location, Save.",
     "Confirm all masters are selectable when creating equipment."],
    "All equipment configuration masters are created successfully.",
    "Medium", "Functional")

add("TC-73", "Equipment", "Create equipment part",
    "TC-72 completed.",
    "QA equipment part",
    ["Open Equipment > Configuration > Equipment Part, click New, enter part name/code and link to the equipment type from TC-72.",
     "Click Save.",
     "Open the Equipment creation form and confirm the part is selectable."],
    "Equipment part is created and is selectable for equipment.",
    "Low", "Functional")

add("TC-74", "Equipment", "Create equipment and make In Service for customer",
    "TC-18/TC-19, TC-72, TC-73 completed.",
    "CUST-QA-001; equipment masters from TC-72/TC-73",
    ["Open Equipment, click New, select Equipment Type/Name/Model/Model No. from TC-72 and Part from TC-73.",
     "Assign the equipment to Customer CUST-QA-001.",
     "Set the equipment Status to 'In Service' and click Save.",
     "Reopen the equipment record and confirm it is linked to CUST-QA-001 with status In Service."],
    "Equipment is associated with customer and status changes to In Service correctly.",
    "High", "Functional")

add("TC-75", "Equipment/Micro Market", "Create equipment for Micro Market with recurring details",
    "TC-72, TC-73 completed; MM-QA-001 master defined.",
    "MM-QA-001; equipment masters from TC-72/TC-73; recurring frequency/amount",
    ["Open Equipment, click New, select equipment masters from TC-72/TC-73.",
     "Assign the equipment to Micro Market MM-QA-001.",
     "Configure recurring billing details (frequency, amount, start date) on the equipment record.",
     "Click Save and reopen the record to confirm the Micro Market association and recurring details are saved."],
    "Equipment is associated with Micro Market and recurring details are saved correctly.",
    "Medium", "Functional")

add("TC-76", "Equipment", "Validate Active/In Service/Repair/Exchange statuses",
    "TC-74 completed.",
    "Equipment record from TC-74",
    ["Open the equipment record from TC-74 and set Status to 'Active'; Save and confirm the status displays correctly.",
     "Change Status to 'In Service'; Save and confirm.",
     "Change Status to 'Repair'; Save and confirm.",
     "Change Status to 'Exchange'; Save and confirm each transition is retained and displayed correctly."],
    "Each equipment status can be selected/updated and displays correctly.",
    "Medium", "Functional")

add("TC-77", "Case Management", "Create case for customer and equipment",
    "TC-71, TC-74 completed.",
    "CUST-QA-001; equipment from TC-74; case category/stage/origin/type from TC-71",
    ["Open Service Management > Cases, click New, select Customer CUST-QA-001 and Equipment from TC-74.",
     "Select Case Category, Stage, Origin and Type created in TC-71, and enter case description.",
     "Click Save and confirm the case is created with correct customer/equipment/configuration links."],
    "Case is created with correct customer, equipment and configuration details.",
    "High", "Functional")

add("TC-78", "Case Management", "Close case with internal and resolution comments",
    "TC-77 completed.",
    "Case from TC-77",
    ["Open the case created in TC-77 and add an Internal Comment.",
     "Add a Resolution Comment describing how the case was resolved.",
     "Change the case Stage/Status to Closed and click Save.",
     "Reopen the case and confirm both comments are retained and the case shows Closed."],
    "Case closes successfully and both comments are retained.",
    "Medium", "Functional")

add("TC-79", "Equipment", "Delete equipment with dependency",
    "TC-74, TC-77 completed (equipment used in a case).",
    "Equipment from TC-74",
    ["Open Equipment, select the record from TC-74 (linked to the TC-77 case) and attempt to delete it.",
     "Confirm the system displays an alert/restriction message and blocks the deletion due to the existing case dependency."],
    "Expected alert/restriction is displayed and unsafe deletion is prevented.",
    "High", "Regression")

add("TC-80", "Reports", "Run all Equipment and Case reports",
    "TC-71 to TC-79 completed.",
    "Equipment/case records from TC-71 to TC-79",
    ["Open Reports > Equipment > Equipment Addition/Archive Report and verify the TC-74/TC-75 equipment appears.",
     "Open Reports > Equipment > Equipment Location Report and verify area/POS location data.",
     "Open Reports > Equipment > Equipment Move Report and verify status changes from TC-76.",
     "Open Reports > Service Management > Case Detailed Report and verify the TC-77/TC-78 case details and comments."],
    "Equipment, case, status and activity reports contain accurate data.",
    "Medium", "Report")

# ---- CRM / Survey / Event (TC-81..86) ----
add("TC-81", "CRM", "Create lead",
    "TC-18/TC-19 completed (or a new prospect available).",
    "CUST-QA-001 or a QA prospect",
    ["Open CRM > Leads, click New, enter prospect/customer name, contact details and source.",
     "Click Save and confirm the lead is created."],
    "Lead is created successfully with required customer/prospect information.",
    "Medium", "Functional")

add("TC-82", "CRM", "Create opportunity",
    "TC-81 completed.",
    "Lead from TC-81",
    ["Open the lead from TC-81 and click Convert to Opportunity (or open CRM > Opportunities > New and link the lead).",
     "Enter opportunity stage, expected revenue and close date.",
     "Click Save and confirm the opportunity is associated with the correct lead/customer."],
    "Opportunity is created and associated correctly.",
    "Medium", "Functional")

add("TC-83", "CRM", "Create CRM event",
    "TC-81/TC-82 completed.",
    "CUST-QA-001; QA event date/participants",
    ["Open CRM > Events, click New, enter event name, date and select participants (CUST-QA-001, EMP-QA-001).",
     "Associate the event with the opportunity/lead from TC-81/TC-82.",
     "Click Save and confirm date, participants and association are correct."],
    "Event is created with correct date, participants and association.",
    "Low", "Functional")

add("TC-84", "Reports", "Run Customer Retention report",
    "TC-22 to TC-49 completed (customer transaction history exists).",
    "CUST-QA-001, CUST-QA-002",
    ["Open Reports > CRM > Customer Retention Report.",
     "Filter by CUST-QA-001 and CUST-QA-002 and run the report.",
     "Verify retention metrics reconcile with the customers' order/invoice history."],
    "Customer retention data is accurate.",
    "Low", "Report")

add("TC-85", "Survey", "Create survey",
    "QA environment available; Survey module access.",
    "QA survey name/questions",
    ["Open Survey app, click New, enter Survey Name and add at least one question/section.",
     "Publish/Save the survey.",
     "Confirm the survey is available for use (e.g. shareable link or assignment)."],
    "Survey is created and available for use.",
    "Low", "Functional")

add("TC-86", "Event", "Create Event",
    "QA environment available; Event module access.",
    "QA event name/date/location",
    ["Open Event module, click New, enter Event Name, date range and location.",
     "Configure required event settings (registration, participants) and Save.",
     "Confirm the event is created successfully with the entered configuration."],
    "Event is created successfully with required configuration.",
    "Low", "Functional")

# ---- Micro Market (TC-87..93) ----
add("TC-87", "Micro Market", "Create Micro Market product catalogue",
    "TC-08 to TC-10 completed.",
    "Catalogue type: Micro Market; PROD-QA-002, PROD-QA-003, PROD-QA-004",
    ["Open Inventory/Micro Market > Product Catalogue, click New and set Catalogue Type to 'Micro Market'.",
     "Add PROD-QA-002, PROD-QA-003 and PROD-QA-004 to the catalogue (excluding PROD-QA-001).",
     "Set price/min/max quantities for each line and click Save."],
    "Catalogue of type Micro Market contains Products 2, 3 and 4.",
    "Medium", "Functional")

add("TC-88", "Micro Market", "Create Micro Market location",
    "TC-18, TC-07, TC-87 completed.",
    "MM-QA-001; CUST-QA-001; PROD-QA-001; catalogue from TC-87",
    ["Open Micro Market > Locations, click New, enter Location Code/Name MM-QA-001.",
     "Link Customer CUST-QA-001 to the location.",
     "Add PROD-QA-001 and the catalogue products from TC-87 to the location's product list.",
     "Click Save."],
    "Location is created with Customer 1, Product 1 and catalogue products.",
    "High", "Functional")

add("TC-89", "Micro Market", "Create store order",
    "TC-88 completed.",
    "MM-QA-001; PROD-QA-001..004; Standard Sales Qty 10",
    ["Open Micro Market > Store Orders, click New, select Location MM-QA-001.",
     "Add PROD-QA-001 with Quantity 10 (and additional catalogue products as needed).",
     "Click Save/Confirm and note the order's stock/order impact."],
    "Store order is created with correct product/quantity and stock/order impact.",
    "Medium", "Functional")

add("TC-90", "Micro Market", "Update location",
    "TC-88 completed.",
    "MM-QA-001",
    ["Open the MM-QA-001 location record and update a field (e.g. address, contact, product list).",
     "Click Save.",
     "Reopen the location and confirm the change is reflected; verify it applies correctly in a subsequent store order/transaction."],
    "Location changes are saved and reflected in subsequent operations.",
    "Low", "Functional")

add("TC-91", "Micro Market", "Create special",
    "TC-88 completed.",
    "MM-QA-001; PROD-QA-002",
    ["Open Micro Market > Specials, click New, select Location MM-QA-001 and product PROD-QA-002.",
     "Configure the special price/discount and effective dates.",
     "Click Save and confirm the special configuration is created successfully."],
    "Special configuration is created successfully.",
    "Low", "Functional")

add("TC-92", "Micro Market", "Create front desk employee discount",
    "TC-01, TC-88 completed.",
    "EMP-QA-001; MM-QA-001; Discount 10%",
    ["Open Micro Market > Front Desk Employee Discount, click New, select EMP-QA-001 and Location MM-QA-001.",
     "Set Discount to 10%.",
     "Click Save and confirm the discount configuration is available for use in a transaction."],
    "Employee discount configuration is created and available for transactions.",
    "Low", "Functional")

add("TC-93", "Micro Market", "Create other payment methods",
    "TC-88 completed.",
    "MM-QA-001; QA payment method name",
    ["Open Micro Market > Configuration > Other Payment Methods, click New, enter a payment method name/code.",
     "Link it to Location MM-QA-001 if applicable and click Save.",
     "Confirm the payment method is selectable during a store transaction."],
    "Other payment methods are configured and available for transaction use.",
    "Low", "Functional")

# ---- Micro Market / Inventory reports (TC-94..101) ----
add("TC-94", "Reports", "Run Stock Analysis report",
    "TC-14, TC-53 to TC-61, TC-89 completed.",
    "PROD-QA-001..004; WH-QA-001",
    ["Open Reports > Inventory > Stock Analysis Report.",
     "Filter by WH-QA-001 and run the report.",
     "Verify current stock levels reconcile with all posted inventory transactions."],
    "Stock analysis reflects current inventory accurately.",
    "Medium", "Report")

add("TC-95", "Reports", "Run Stock Value report",
    "TC-14, TC-53 to TC-61 completed.",
    "PROD-QA-001..004; WH-QA-001",
    ["Open Reports > Inventory > Stock Value Report.",
     "Run the report for WH-QA-001 and verify values reconcile with quantities and costs recorded in TC-14/TC-53 to TC-61."],
    "Stock values reconcile with inventory quantities and costs.",
    "Medium", "Report")

add("TC-96", "Reports", "Run Store Order List report",
    "TC-89 completed.",
    "MM-QA-001",
    ["Open Reports > Micro Market > Store Order List Report.",
     "Filter by MM-QA-001 and run the report.",
     "Verify the TC-89 store order and its status are listed accurately."],
    "Store orders and statuses are accurate.",
    "Medium", "Report")

add("TC-97", "Reports", "Run Product Move report",
    "TC-69, TC-70, TC-89 completed.",
    "PROD-QA-001..004",
    ["Open Reports > Inventory > Product Move Report.",
     "Filter by PROD-QA-001 and run the report.",
     "Verify all product movements (transfer, store order, adjustment) are accurately listed."],
    "Product movements are accurately listed.",
    "Medium", "Report")

add("TC-98", "Reports", "Run Subsidy report",
    "TC-18 (Customer 1 subsidy configured), TC-22/TC-24 transactions completed.",
    "CUST-QA-001",
    ["Open Reports > Micro Market/Customer > Subsidy Report.",
     "Filter by CUST-QA-001 and run the report.",
     "Verify subsidy amounts and transactions reconcile with the configured subsidy on CUST-QA-001."],
    "Subsidy amounts and transactions are accurate.",
    "Medium", "Report")

add("TC-99", "Reports", "Run Unpicked Product List report",
    "TC-25, TC-44 completed (orders with pending delivery/picking).",
    "PROD-QA-001",
    ["Open Reports > Inventory/Sales > Unpicked Product List Report.",
     "Run the report and verify any product/order lines pending pick are correctly identified (e.g. the TC-44 backorder remainder before delivery)."],
    "Unpicked products are accurately identified.",
    "Low", "Report")

add("TC-100", "Reports", "Run Where Sold report",
    "TC-22 to TC-44 completed.",
    "PROD-QA-001..004",
    ["Open Reports > Sales > Where Sold Report.",
     "Filter by PROD-QA-001 to PROD-QA-004 and run the report.",
     "Verify sales locations/customers are accurately displayed for each product."],
    "Product sales locations/customers are accurately displayed.",
    "Medium", "Report")

add("TC-101", "Reports", "Run Products with Min/Max report",
    "TC-14, TC-53 to TC-61 completed; min/max levels configured on a product.",
    "PROD-QA-001 (min/max configured)",
    ["Configure Minimum/Maximum stock levels on PROD-QA-001 (Inventory > Products > Reordering rules).",
     "Adjust stock via TC-14/TC-53 to bring the quantity below minimum or above maximum.",
     "Open Reports > Inventory > Products with Min/Max Report and run it.",
     "Verify PROD-QA-001 is correctly flagged as at/below minimum or at/above maximum."],
    "Products at/below minimum or at/above maximum are correctly reported.",
    "Medium", "Report")

# ---- Transactions / variance reports (TC-102..116) ----
add("TC-102", "Transactions", "Perform a transaction and validate transaction-related reports",
    "TC-89 or TC-22 completed (a sales/store transaction exists).",
    "MM-QA-001 or CUST-QA-001; PROD-QA-001",
    ["Perform a store/sales transaction (e.g. the TC-89 store order or a POS sale) for PROD-QA-001.",
     "Open Reports > Transactions > Transaction Report, filter by the transaction date and verify it is recorded correctly.",
     "Open Reports > Sales > Sales by Product Report and verify the transaction is reflected for PROD-QA-001.",
     "Open Reports > Disbursement Report and verify the disbursement entry matches the transaction."],
    "Transaction is recorded and reflected correctly in Transaction, Sales by Product and Disbursement reports.",
    "High", "Integration")

add("TC-103", "Inventory/Reports", "Perform inventory transaction and validate Overage/Spoilage/Shrinkage",
    "TC-14 completed (adjustment capability available).",
    "PROD-QA-001; Quantity Adjustment (overage/shrinkage) values",
    ["Open Inventory > Inventory Adjustments, create an adjustment on PROD-QA-001 representing an overage (positive variance) and Validate.",
     "Create a second adjustment representing a shrinkage/spoilage (negative variance) and Validate.",
     "Open Reports > Inventory > Overage Report and verify the overage entry.",
     "Open Reports > Inventory > Spoilage Report and Shrinkage Report and verify the corresponding entries."],
    "Inventory adjustment results are reflected correctly in Overage, Spoilage and Shrinkage reports.",
    "Medium", "Report")

add("TC-104", "Reports", "Run warehouse Overage/Spoilage/Shrinkage report",
    "TC-103 completed.",
    "WH-QA-001",
    ["Open Reports > Inventory > Overage/Spoilage/Shrinkage Report (Warehouse).",
     "Filter by WH-QA-001 and run the report.",
     "Verify warehouse-specific variance values match the TC-103 adjustments."],
    "Warehouse-specific variance values are accurate.",
    "Medium", "Report")

add("TC-105", "Reports", "Run Inventory Change report",
    "TC-14, TC-53 to TC-61, TC-69, TC-70, TC-103 completed.",
    "PROD-QA-001..004",
    ["Open Reports > Inventory > Inventory Change Report.",
     "Run the report for the QA date range and verify all inventory-affecting transactions are accurately reported."],
    "Inventory changes are accurately reported.",
    "Medium", "Report")

add("TC-106", "Reports", "Run Inventory Change (Warehouse) report",
    "TC-105 completed.",
    "WH-QA-001",
    ["Open Reports > Inventory > Inventory Change Report (Warehouse).",
     "Filter by WH-QA-001 and run the report.",
     "Verify warehouse-specific inventory changes are accurate and reconcile with TC-105."],
    "Warehouse-specific inventory changes are accurate.",
    "Medium", "Report")

add("TC-107", "Reports", "Run Customer Commission report",
    "TC-22 to TC-44 completed; commission configured for a customer/employee.",
    "CUST-QA-001; EMP-QA-001",
    ["Open Reports > Sales > Customer Commission Report.",
     "Filter by CUST-QA-001/EMP-QA-001 and run the report.",
     "Verify commission values reconcile with the eligible transactions from TC-22 to TC-44."],
    "Customer commission values reconcile with eligible transactions.",
    "Low", "Report")

add("TC-108", "Reports", "Run Employee Discount report",
    "TC-92 completed (front desk employee discount used in a transaction).",
    "EMP-QA-001; MM-QA-001",
    ["Perform a Micro Market transaction using the TC-92 front desk employee discount.",
     "Open Reports > Micro Market > Employee Discount Report.",
     "Filter by EMP-QA-001 and run the report; verify the discount transaction and value are accurate."],
    "Employee discount transactions and values are accurate.",
    "Low", "Report")

add("TC-109", "Reports", "Run Payroll Deduction report",
    "TC-108 completed (employee discount/deduction transaction exists).",
    "EMP-QA-001",
    ["Open Reports > Payroll > Payroll Deduction Report.",
     "Filter by EMP-QA-001 and run the report.",
     "Verify payroll deduction values reconcile with the applicable employee transactions from TC-108."],
    "Payroll deduction values reconcile with applicable employee transactions.",
    "Low", "Report")

add("TC-110", "Reports", "Run MM Sales Analysis report",
    "TC-89, TC-102 completed.",
    "MM-QA-001",
    ["Open Reports > Micro Market > MM Sales Analysis Report.",
     "Filter by MM-QA-001 and run the report.",
     "Verify sales analysis figures accurately reflect the store order/transaction activity from TC-89/TC-102."],
    "Micro Market sales analysis accurately reflects transactions.",
    "Medium", "Report")

add("TC-111", "Reports", "Run Periodic Sales report",
    "TC-89, TC-102 completed.",
    "MM-QA-001",
    ["Open Reports > Micro Market > Periodic Sales Report.",
     "Run the report for the QA period and verify totals reconcile with underlying transactions."],
    "Periodic sales totals reconcile with underlying transactions.",
    "Low", "Report")

add("TC-112", "Reports", "Run Weekly Sales report",
    "TC-89, TC-102 completed.",
    "MM-QA-001",
    ["Open Reports > Micro Market > Weekly Sales Report.",
     "Run the report for the current week and verify sales totals and period boundaries reconcile with transaction data."],
    "Weekly sales totals and periods reconcile with transaction data.",
    "Low", "Report")

add("TC-113", "Reports", "Run Product Rank report",
    "TC-22 to TC-44, TC-89 completed.",
    "PROD-QA-001..004",
    ["Open Reports > Sales > Product Rank Report.",
     "Run the report and verify products are ranked by the correct sales/quantity criteria based on QA transactions."],
    "Product ranking is based on the correct sales/quantity criteria.",
    "Low", "Report")

add("TC-114", "Reports", "Run Terminal report",
    "TC-89, TC-102 completed (POS/terminal transaction exists).",
    "MM-QA-001",
    ["Open Reports > Micro Market > Terminal Report.",
     "Filter by MM-QA-001/terminal and run the report.",
     "Verify terminal transactions and totals are accurately reported."],
    "Terminal transactions and totals are accurately reported.",
    "Low", "Report")

add("TC-115", "Reports", "Run Daily Net Sales by Customer report",
    "TC-22 to TC-44 completed.",
    "CUST-QA-001, CUST-QA-002",
    ["Open Reports > Sales > Daily Net Sales by Customer Report.",
     "Filter by today's date and the QA customers, run the report.",
     "Verify daily net sales reconcile with invoices/payments/transactions posted for CUST-QA-001 and CUST-QA-002."],
    "Daily net sales by customer reconcile with invoices/payments/transactions.",
    "Medium", "Report")

add("TC-116", "Reports", "Run Monthly Net Sales by Customer report",
    "TC-115 completed.",
    "CUST-QA-001, CUST-QA-002",
    ["Open Reports > Sales > Monthly Net Sales by Customer Report.",
     "Filter by the current month and the QA customers, run the report.",
     "Verify monthly net sales reconcile with the underlying sales data from TC-22 to TC-44."],
    "Monthly net sales by customer reconcile with underlying sales data.",
    "Medium", "Report")

# ---- Regression (TC-117, TC-118) ----
add("TC-117", "Regression", "Attempt to delete product used in transaction",
    "TC-22 to TC-116 completed (product has extensive transaction history).",
    "PROD-QA-001",
    ["Open Inventory > Products, select PROD-QA-001 (used across sales, purchase, transfer and report transactions).",
     "Attempt to delete PROD-QA-001.",
     "Confirm the system displays the expected dependency alert/restriction message and blocks the deletion.",
     "Verify all historical transactions referencing PROD-QA-001 remain intact and unaffected."],
    "System displays the expected dependency alert/restriction and preserves historical transaction integrity.",
    "High", "Regression")

add("TC-118", "Regression/Reports", "Rerun historical transaction and sales reports after product deletion attempt",
    "TC-117 completed.",
    "PROD-QA-001; all prior transaction/report data",
    ["Re-run Reports > Transactions > Transaction Report and verify prior PROD-QA-001 entries are unchanged.",
     "Re-run Sales by Product and Disbursement reports and verify totals are unchanged.",
     "Re-run Customer Commission, Employee Discount and Payroll Deduction reports and verify values are unchanged.",
     "Re-run MM Sales Analysis, Periodic Sales, Weekly Sales, Product Rank and Terminal reports and verify values are unchanged.",
     "Re-run Daily Net Sales by Customer and Monthly Net Sales by Customer reports and verify values are unchanged."],
    "Transaction, Sales by Product, Disbursement, Customer Commission, Employee Discount, Payroll Deduction, MM Sales Analysis, Periodic Sales, Weekly Sales, Product Rank, Terminal, Daily Net Sales by Customer and Monthly Net Sales by Customer remain accurate.",
    "High", "Regression")

assert len(TC) == 118, f"Expected 118 test cases, got {len(TC)}"

# ---------------------------------------------------------------------------
# Build workbook
# ---------------------------------------------------------------------------
wb = openpyxl.Workbook()

HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
WRAP = Alignment(wrap_text=True, vertical="top")
WRAP_CENTER = Alignment(wrap_text=True, vertical="top", horizontal="center")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

def style_header(ws, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = WRAP_CENTER
        cell.border = BORDER
    ws.freeze_panes = "A2"

# --- Sheet 1: Test Cases ---
ws = wb.active
ws.title = "Test Cases"
headers = ["TC ID", "Module", "Test Scenario", "Precondition", "Test Data", "Test Steps", "Expected Result", "Priority", "Type"]
ws.append(headers)
widths = [9, 16, 30, 26, 26, 55, 40, 10, 12]
for i, w in enumerate(widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = w

for tc in TC:
    row = [tc["id"], tc["module"], tc["scenario"], tc["precondition"], tc["data"], tc["steps"], tc["expected"], tc["priority"], tc["type"]]
    ws.append(row)

for r in range(2, len(TC) + 2):
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=r, column=c)
        cell.alignment = WRAP
        cell.border = BORDER
    ws.row_dimensions[r].height = 90

style_header(ws, len(headers))

# --- Sheet 2: Test Data (master + transaction data, from plan sections 5 & 6) ---
ws2 = wb.create_sheet("Test Data")
ws2.append(["Data Type", "Code", "Description"])
for row in MASTER_DATA:
    ws2.append(list(row))
ws2.append([])
ws2.append(["Transaction Data", "Value", ""])
for row in TXN_DATA:
    ws2.append([row[0], row[1], ""])
for i, w in enumerate([22, 30, 45], 1):
    ws2.column_dimensions[get_column_letter(i)].width = w
for r in range(2, ws2.max_row + 1):
    for c in range(1, 4):
        cell = ws2.cell(row=r, column=c)
        cell.alignment = WRAP
        cell.border = BORDER
style_header(ws2, 3)

# --- Sheet 3: Test Summary ---
ws3 = wb.create_sheet("Test Summary")
ws3.append(["Metric", "Value"])
from collections import Counter
mod_counts = Counter(tc["module"] for tc in TC)
type_counts = Counter(tc["type"] for tc in TC)
prio_counts = Counter(tc["priority"] for tc in TC)
ws3.append(["Total Test Cases", len(TC)])
for t, c in sorted(type_counts.items()):
    ws3.append([f"Type: {t}", c])
for p in ["High", "Medium", "Low"]:
    ws3.append([f"Priority: {p}", prio_counts.get(p, 0)])
ws3.append([])
ws3.append(["Module", "Test Case Count"])
for m, c in sorted(mod_counts.items()):
    ws3.append([m, c])
ws3.column_dimensions["A"].width = 30
ws3.column_dimensions["B"].width = 16
for r in range(2, ws3.max_row + 1):
    for c in range(1, 3):
        ws3.cell(row=r, column=c).border = BORDER
style_header(ws3, 2)

# --- Sheet 4: Execution ---
ws4 = wb.create_sheet("Execution")
ws4.append(["TC ID", "Module", "Test Scenario", "Actual Result", "Status", "Defect ID", "Tester", "Execution Date", "Remarks"])
for tc in TC:
    ws4.append([tc["id"], tc["module"], tc["scenario"], "", "Not Executed", "", "", "", ""])
widths4 = [9, 16, 30, 30, 14, 12, 14, 16, 26]
for i, w in enumerate(widths4, 1):
    ws4.column_dimensions[get_column_letter(i)].width = w
for r in range(2, ws4.max_row + 1):
    for c in range(1, 10):
        cell = ws4.cell(row=r, column=c)
        cell.alignment = WRAP
        cell.border = BORDER
style_header(ws4, 9)

out_path = "/home/shankar/playwright-odoo/Averigo_End_to_End_Test_Cases.xlsx"
wb.save(out_path)
print("Saved:", out_path)
print("Total test cases:", len(TC))
